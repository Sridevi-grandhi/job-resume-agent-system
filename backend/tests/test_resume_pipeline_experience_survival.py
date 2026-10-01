"""End-to-end regression: Professional Experience (multiple employers, a mix
of marker-based and marker-less bullets) must survive intact through the
FULL real pipeline — original upload text extraction, structure parsing,
LLM-result assembly, the Step 5 editor's structured document model, a
save round-trip (re-serializing the editor document back to plain text,
exactly like a real "Save" click), the final DOCX render, and (where
LibreOffice is available) the exported PDF.

This closes a real gap in the existing regression suite: prior DOCX-level
tests only asserted heading occurrence counts, never that actual employer/
bullet body text survives. Every fixture uses a different, parametrized
fake identity — never a real person's data — to prove this is a general
fix, not a patch for one candidate's resume shape.
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.shared import RGBColor

from app.services.resume_assembly_service import assemble_optimized_resume
from app.services.resume_document_service import document_to_plain_text, plain_text_to_document
from app.services.resume_structure_parser import parse_resume_structure
from app.services.template_preservation_service import render_optimized_document, validate_rendered_structure
from app.utils.resume_parser import parse_docx

BLUE = (0x1F, 0x4E, 0x79)

SKILLS_ROWS = [
    ("Languages", "Java, Python, Go, C++, Rust, Kotlin"),
    ("Frameworks", "Spring Boot, Django, gRPC, FastAPI, Netty"),
    ("Databases", "PostgreSQL, MongoDB, Redis, Cassandra, DynamoDB"),
    ("Cloud", "AWS, GCP, Azure, DigitalOcean"),
    ("DevOps", "Docker, Kubernetes, Terraform, Jenkins, ArgoCD"),
    ("Testing", "JUnit, PyTest, Selenium, Cypress, k6"),
    ("Messaging", "Kafka, RabbitMQ, SQS, Pulsar"),
    ("Monitoring", "Prometheus, Grafana, Splunk, Datadog"),
    ("Scripting", "Bash, PowerShell, Groovy, Perl"),
    ("Build Tools", "Maven, Gradle, Webpack, Bazel"),
    ("Version Control", "Git, GitHub Actions, GitLab CI"),
    ("Security", "OAuth2, JWT, Vault, SIEM"),
    ("Collaboration", "Jira, Confluence, Slack"),
]

EMPLOYERS = [
    {
        "company": "CloudScale Inc",
        "location": "Denver, CO",
        "title": "Senior Backend Engineer",
        "dates": "Jun 2021 - Present",
        # Marker-based bullets — the common, already-supported case.
        "bullets": [
            "- Designed microservices handling millions of daily requests.",
            "- Migrated the payments pipeline to an event-driven architecture.",
        ],
    },
    {
        "company": "Northlight Analytics",
        "location": "Boston, MA",
        "title": "Backend Engineer",
        "dates": "Jan 2019 - May 2021",
        # Marker-less bullets — exercises the widened bullet recognition.
        "bullets": [
            "Collaborated with cross-functional teams to redesign the checkout flow and improve conversion.",
            "Reduced infrastructure costs by consolidating redundant services across environments.",
        ],
    },
    {
        "company": "Riverbend Software",
        "location": "Austin, TX",
        "title": "Software Engineer",
        "dates": "Aug 2016 - Dec 2018",
        # Mixed marker / marker-less bullets under the same employer.
        "bullets": [
            "- Built internal tooling used by every engineering team.",
            "Automated release pipelines that cut deployment time by half.",
        ],
    },
]


def _add_paragraph(document, text, *, bold=False, italic=False, color=None):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return paragraph


def _bullet_text(line: str) -> str:
    return line[2:] if line.startswith("- ") else line


def _header_line(employer):
    # Company + location + dates on one line — a DATE_PATTERN match on the
    # header line itself is the simplest signal _parse_employment_blocks
    # uses to mark "a new employer starts here". The split-header variant
    # (company/location/role/dates each on their own line, no date on the
    # header line at all) is covered separately below.
    return f"{employer['company']} | {employer['location']}    {employer['dates']}"


def _split_header_lines(employer):
    # Company, location, role, and dates each on their own line — no date
    # on the "header" line itself. _parse_employment_blocks recognizes this
    # via a conservative look-ahead (a company-shaped line with a date a
    # few lines below it, nothing bullet-like in between) rather than the
    # simple same-line DATE_PATTERN match _header_line relies on.
    return [employer["company"], employer["location"], employer["title"], employer["dates"]]


def _build_docx(path, *, name, contact, experience_heading="PROFESSIONAL EXPERIENCE"):
    document = Document()
    _add_paragraph(document, name)
    _add_paragraph(document, contact)
    _add_paragraph(document, "PROFESSIONAL SUMMARY")
    _add_paragraph(document, "Backend engineer.")
    _add_paragraph(document, "TECHNICAL SKILLS")
    table = document.add_table(rows=len(SKILLS_ROWS), cols=2)
    for row, (label, values) in zip(table.rows, SKILLS_ROWS):
        row.cells[0].text = label
        row.cells[1].text = values
    document.add_paragraph()  # blank paragraph before Experience, per the reported shape
    _add_paragraph(document, experience_heading)
    for employer in EMPLOYERS:
        _add_paragraph(document, _header_line(employer), bold=True, color=BLUE)
        _add_paragraph(document, f"Role: {employer['title']}", bold=True, color=BLUE)
        for line in employer["bullets"]:
            _add_paragraph(document, line)
    document.save(path)


def _optimized_text_for(name, contact, experience_heading):
    skills_lines = "\n".join(f"{label}: {values}" for label, values in SKILLS_ROWS)
    experience_lines = []
    for employer in EMPLOYERS:
        experience_lines.append(_header_line(employer))
        experience_lines.append(f"Role: {employer['title']}")
        experience_lines.extend(employer["bullets"])
        experience_lines.append("")
    return (
        f"{name}\n{contact}\n\n"
        f"Professional Summary\nBackend engineer.\n\n"
        f"Technical Skills\n{skills_lines}\n\n"
        f"{experience_heading.title()}\n" + "\n".join(experience_lines)
    )


def _build_docx_split_header(path, *, name, contact, experience_heading="PROFESSIONAL EXPERIENCE"):
    document = Document()
    _add_paragraph(document, name)
    _add_paragraph(document, contact)
    _add_paragraph(document, "PROFESSIONAL SUMMARY")
    _add_paragraph(document, "Backend engineer.")
    _add_paragraph(document, "TECHNICAL SKILLS")
    table = document.add_table(rows=len(SKILLS_ROWS), cols=2)
    for row, (label, values) in zip(table.rows, SKILLS_ROWS):
        row.cells[0].text = label
        row.cells[1].text = values
    document.add_paragraph()
    _add_paragraph(document, experience_heading)
    for employer in EMPLOYERS:
        for line in _split_header_lines(employer):
            _add_paragraph(document, line, bold=True, color=BLUE)
        for line in employer["bullets"]:
            _add_paragraph(document, line)
    document.save(path)


def _optimized_text_for_split_header(name, contact, experience_heading):
    skills_lines = "\n".join(f"{label}: {values}" for label, values in SKILLS_ROWS)
    experience_lines = []
    for employer in EMPLOYERS:
        experience_lines.extend(_split_header_lines(employer))
        experience_lines.extend(employer["bullets"])
        experience_lines.append("")
    return (
        f"{name}\n{contact}\n\n"
        f"Professional Summary\nBackend engineer.\n\n"
        f"Technical Skills\n{skills_lines}\n\n"
        f"{experience_heading.title()}\n" + "\n".join(experience_lines)
    )


@pytest.mark.parametrize(
    "name,contact,experience_heading",
    [
        ("Priya Patel", "priya.patel@example.com | 555-222-3333", "PROFESSIONAL EXPERIENCE"),
        ("Marcus Webb", "marcus.webb@example.org | 444-777-9999", "WORK HISTORY"),
    ],
    ids=["standard-heading", "work-history-heading"],
)
def test_all_employers_survive_the_full_pipeline(tmp_path, name, contact, experience_heading):
    source = tmp_path / "source.docx"
    _build_docx(source, name=name, contact=contact, experience_heading=experience_heading)

    # 1) Upload-time text extraction — real resume_parser.parse_docx(), not a stub.
    resume_text = parse_docx(str(source))
    structure = parse_resume_structure(resume_text)
    assert len(structure["employments"]) == len(EMPLOYERS), (
        f"original structure lost an employer: {[e.get('company') for e in structure['employments']]}"
    )

    # 2) LLM result — assembly must fall back to the ORIGINAL experience content
    # untouched (no rewritten_bullets supplied), matching a real "GPT omitted
    # experience updates" response.
    llm_result = {
        "generated_resume_plain_text": _optimized_text_for(name, contact, experience_heading),
    }
    assembled = assemble_optimized_resume(resume_text=resume_text, structure=structure, llm_result=llm_result)
    optimized_text = assembled["plain_text"]
    for employer in EMPLOYERS:
        assert employer["company"] in optimized_text
        for line in employer["bullets"]:
            assert _bullet_text(line) in optimized_text

    # 3) Step 5 editor's structured document model.
    document = plain_text_to_document(optimized_text)
    experience_sections = [s for s in document["sections"] if s["type"] == "experience"]
    assert experience_sections, "no experience section reached the editor document at all"
    editor_companies = [item["company"] for section in experience_sections for item in section["items"]]
    for employer in EMPLOYERS:
        assert employer["company"] in editor_companies, f"{employer['company']} missing from editor document"
    editor_bullets = [
        b["text"]
        for section in experience_sections
        for item in section["items"]
        for b in item["bullets"]
    ]
    for employer in EMPLOYERS:
        for line in employer["bullets"]:
            assert _bullet_text(line) in editor_bullets, f"bullet lost in editor document: {line!r}"

    # 4) Save round-trip — re-serialize the editor document back to plain text,
    # exactly like a real "Save" click (optimizer_edit_routes._resolve_edited_payload).
    saved_text = document_to_plain_text(document)
    for employer in EMPLOYERS:
        assert employer["company"] in saved_text, f"{employer['company']} lost on save round-trip"
        for line in employer["bullets"]:
            assert _bullet_text(line) in saved_text, f"bullet lost on save round-trip: {line!r}"

    # 5) Final DOCX render.
    rendered = render_optimized_document(str(source), "docx", saved_text, output_path=str(tmp_path / "out.docx"))
    result_doc = Document(rendered["file_path"])
    body_text = "\n".join(p.text for p in result_doc.paragraphs)
    for employer in EMPLOYERS:
        assert employer["company"] in body_text, f"{employer['company']} missing from final DOCX"
        for line in employer["bullets"]:
            assert _bullet_text(line) in body_text, f"bullet missing from final DOCX: {line!r}"

    report = validate_rendered_structure(
        result_doc, saved_text, original_employer_names=[e["company"] for e in EMPLOYERS],
    )
    assert report["duplicate_section_kinds"] == []
    assert report["employer_count_found"] == report["employer_count_expected"] == len(EMPLOYERS)
    assert report["missing_employers"] == []


def test_split_header_employers_survive_the_full_pipeline(tmp_path):
    """Same three employers, same full pipeline, but each header spread
    across 4 separate lines (company / location / role / dates) with no
    date on the header line itself — the split-header shape Gap 1 fixes."""
    name, contact = "Dana Whitfield", "dana.whitfield@example.com | 555-333-1111"
    source = tmp_path / "source.docx"
    _build_docx_split_header(source, name=name, contact=contact)

    resume_text = parse_docx(str(source))
    structure = parse_resume_structure(resume_text)
    assert len(structure["employments"]) == len(EMPLOYERS), (
        f"original structure lost an employer: {[e.get('company') for e in structure['employments']]}"
    )
    for employment, employer in zip(structure["employments"], EMPLOYERS):
        assert employment["company"] == employer["company"]
        assert employment["dates"] == employer["dates"]

    llm_result = {
        "generated_resume_plain_text": _optimized_text_for_split_header(name, contact, "PROFESSIONAL EXPERIENCE"),
    }
    assembled = assemble_optimized_resume(resume_text=resume_text, structure=structure, llm_result=llm_result)
    optimized_text = assembled["plain_text"]
    for employer in EMPLOYERS:
        assert employer["company"] in optimized_text
        for line in employer["bullets"]:
            assert _bullet_text(line) in optimized_text

    document = plain_text_to_document(optimized_text)
    experience_sections = [s for s in document["sections"] if s["type"] == "experience"]
    editor_companies = [item["company"] for section in experience_sections for item in section["items"]]
    for employer in EMPLOYERS:
        assert employer["company"] in editor_companies, f"{employer['company']} missing from editor document"
    editor_bullets = [b["text"] for section in experience_sections for item in section["items"] for b in item["bullets"]]
    for employer in EMPLOYERS:
        for line in employer["bullets"]:
            assert _bullet_text(line) in editor_bullets, f"bullet lost in editor document: {line!r}"

    saved_text = document_to_plain_text(document)
    for employer in EMPLOYERS:
        assert employer["company"] in saved_text
        for line in employer["bullets"]:
            assert _bullet_text(line) in saved_text

    rendered = render_optimized_document(str(source), "docx", saved_text, output_path=str(tmp_path / "out-split.docx"))
    result_doc = Document(rendered["file_path"])
    body_text = "\n".join(p.text for p in result_doc.paragraphs)
    for employer in EMPLOYERS:
        assert employer["company"] in body_text, f"{employer['company']} missing from final DOCX"
        for line in employer["bullets"]:
            assert _bullet_text(line) in body_text, f"bullet missing from final DOCX: {line!r}"

    report = validate_rendered_structure(
        result_doc, saved_text, original_employer_names=[e["company"] for e in EMPLOYERS],
    )
    assert report["duplicate_section_kinds"] == []
    assert report["employer_count_found"] == report["employer_count_expected"] == len(EMPLOYERS)
    assert report["missing_employers"] == []


def _soffice_available():
    from app.services.docx_pdf_export import _find_soffice
    return _find_soffice() is not None


@pytest.mark.skipif(not _soffice_available(), reason="LibreOffice (soffice) not installed locally")
def test_all_employers_survive_into_the_exported_pdf(tmp_path):
    import fitz

    from app.services.docx_pdf_export import convert_docx_to_pdf, re_sub_ws

    name, contact = "Priya Patel", "priya.patel@example.com | 555-222-3333"
    source = tmp_path / "source.docx"
    _build_docx(source, name=name, contact=contact)

    resume_text = parse_docx(str(source))
    structure = parse_resume_structure(resume_text)
    llm_result = {"generated_resume_plain_text": _optimized_text_for(name, contact, "PROFESSIONAL EXPERIENCE")}
    assembled = assemble_optimized_resume(resume_text=resume_text, structure=structure, llm_result=llm_result)
    document = plain_text_to_document(assembled["plain_text"])
    saved_text = document_to_plain_text(document)

    rendered = render_optimized_document(str(source), "docx", saved_text, output_path=str(tmp_path / "out.docx"))
    pdf_path = convert_docx_to_pdf(rendered["file_path"], tmp_path / "out.pdf")
    pdf = fitz.open(pdf_path)
    try:
        text = re_sub_ws("\n".join(page.get_text("text") or "" for page in pdf))
    finally:
        pdf.close()

    for employer in EMPLOYERS:
        assert employer["company"] in text, f"{employer['company']} missing from exported PDF"


@pytest.mark.skipif(not _soffice_available(), reason="LibreOffice (soffice) not installed locally")
def test_split_header_employers_survive_into_the_exported_pdf(tmp_path):
    import fitz

    from app.services.docx_pdf_export import convert_docx_to_pdf, re_sub_ws

    name, contact = "Dana Whitfield", "dana.whitfield@example.com | 555-333-1111"
    source = tmp_path / "source.docx"
    _build_docx_split_header(source, name=name, contact=contact)

    resume_text = parse_docx(str(source))
    structure = parse_resume_structure(resume_text)
    llm_result = {
        "generated_resume_plain_text": _optimized_text_for_split_header(name, contact, "PROFESSIONAL EXPERIENCE"),
    }
    assembled = assemble_optimized_resume(resume_text=resume_text, structure=structure, llm_result=llm_result)
    document = plain_text_to_document(assembled["plain_text"])
    saved_text = document_to_plain_text(document)

    rendered = render_optimized_document(str(source), "docx", saved_text, output_path=str(tmp_path / "out-split.docx"))
    pdf_path = convert_docx_to_pdf(rendered["file_path"], tmp_path / "out-split.pdf")
    pdf = fitz.open(pdf_path)
    try:
        text = re_sub_ws("\n".join(page.get_text("text") or "" for page in pdf))
    finally:
        pdf.close()

    for employer in EMPLOYERS:
        assert employer["company"] in text, f"{employer['company']} missing from exported PDF"
        assert employer["dates"] in text, f"{employer['dates']} missing from exported PDF for {employer['company']}"
