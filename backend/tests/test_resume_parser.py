"""Tests for app/utils/resume_parser.py."""

from __future__ import annotations

import pytest
from docx import Document

from app.utils.resume_parser import parse_docx, parse_pdf, parse_resume_file, parse_txt


def test_parse_txt_reads_and_strips_content(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("  Java developer with Spring Boot experience.  \n", encoding="utf-8")
    assert parse_txt(str(path)) == "Java developer with Spring Boot experience."


def test_parse_txt_ignores_undecodable_bytes(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_bytes("Java developer \xff\xfe garbled".encode("latin-1"))
    result = parse_txt(str(path))
    assert "Java developer" in result


def test_parse_docx_joins_nonblank_paragraphs(tmp_path):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Developer")
    document.add_paragraph("")
    document.add_paragraph("Java Spring Boot experience.")
    document.save(str(path))

    result = parse_docx(str(path))
    assert result == "Jane Developer\nJava Spring Boot experience."


def test_parse_pdf_extracts_text(tmp_path):
    import fitz

    path = tmp_path / "resume.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Jane Developer - Java Engineer")
    doc.save(str(path))
    doc.close()

    result = parse_pdf(str(path))
    assert "Jane Developer" in result


def test_parse_resume_file_dispatches_by_extension_txt(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text("Java developer", encoding="utf-8")
    assert parse_resume_file(str(path)) == "Java developer"


def test_parse_resume_file_dispatches_by_extension_docx(tmp_path):
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Java developer")
    document.save(str(path))
    assert parse_resume_file(str(path)) == "Java developer"


def test_parse_resume_file_unsupported_extension_raises_valueerror(tmp_path):
    path = tmp_path / "resume.exe"
    path.write_bytes(b"binary")
    with pytest.raises(ValueError, match="Unsupported file type"):
        parse_resume_file(str(path))


def test_parse_resume_file_extension_is_case_insensitive(tmp_path):
    path = tmp_path / "resume.TXT"
    path.write_text("Java developer", encoding="utf-8")
    assert parse_resume_file(str(path)) == "Java developer"
