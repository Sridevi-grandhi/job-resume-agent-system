from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.shared import Inches, Pt, RGBColor

from app.models.optimized_resume_model import OptimizedResume
from app.services.ats_optimization_workflow_service import optimization_eligibility, role_aware_optimization_eligibility
from app.services.ats_scoring import calculate_weighted_ats_score, classify_role, classify_role_families
from app.services.requirement_extraction_service import group_canonical_score_evidence
from app.services.template_preservation_service import render_optimized_document


@pytest.mark.parametrize("score,allowed", [(64.9, False), (65.0, True), (65.1, True)])
def test_eligibility_threshold_is_exact(score, allowed):
    result = optimization_eligibility(score)
    assert result["optimization_allowed"] is allowed
    assert result["minimum_required_score"] == 65.0


def test_eligibility_tiers():
    assert optimization_eligibility(30)["eligibility_tier"] == "NOT_RECOMMENDED"
    assert optimization_eligibility(50)["eligibility_tier"] == "LOW_ALIGNMENT"
    assert optimization_eligibility(70)["eligibility_tier"] == "ELIGIBLE_FOR_OPTIMIZATION"
    assert optimization_eligibility(80)["eligibility_tier"] == "STRONG_OPTIMIZATION_CANDIDATE"
    assert optimization_eligibility(95)["eligibility_tier"] == "ALREADY_HIGH_MATCH"


@pytest.mark.parametrize(
    "resume_text,job_title,expected",
    [
        ("Senior DevOps Engineer\nAWS Docker Kubernetes Terraform Jenkins", "Senior DevOps Engineer", "EXACT"),
        ("DevOps Engineer", "Senior DevOps Engineer / Salt Lake City / AWS", "EXACT"),
        ("DevOps Engineer", "Site Reliability Engineer", "STRONG_RELATED"),
        ("DevOps Engineer", "Platform Engineer", "STRONG_RELATED"),
        ("DevOps Engineer", "Java Full Stack Developer", "MISMATCH"),
    ],
)
def test_canonical_role_alignment(resume_text, job_title, expected):
    job = SimpleNamespace(title=job_title, description=f"Hiring a {job_title}", skills_required="", location="", work_mode="", job_type="")
    result = calculate_weighted_ats_score(resume_text, job)
    assert result["role_alignment"] == expected


def test_low_score_is_separate_from_exact_role_alignment():
    job = SimpleNamespace(
        title="Senior DevOps Engineer", description="Requires Kubernetes Terraform Argo CD and Helm",
        skills_required="Kubernetes, Terraform, Argo CD, Helm", location="", work_mode="", job_type="",
    )
    result = calculate_weighted_ats_score("DevOps Engineer\nAWS Docker Jenkins", job)
    assert result["role_alignment"] == "EXACT"
    assert result["match_percentage"] < 65
    assert optimization_eligibility(result["match_percentage"])["optimization_allowed"] is False


def test_preferences_are_not_candidate_role_evidence():
    # The classifier accepts resume content and a document title only; no preference
    # input exists, so a DevOps preference cannot overwrite Python resume evidence.
    assert classify_role_families("Python Developer\nBuilt APIs with Python and FastAPI") == ["Python Backend"]


@pytest.mark.parametrize(
    "resume,job_title,expected",
    [
        ("Java Full Stack Developer using Spring Boot and React", "Spring Boot Developer", "STRONG_RELATED"),
        ("Senior SDET building Selenium automation", "QA Automation Engineer", "EXACT"),
        ("Data Engineer building Spark ETL pipelines", "Data Platform Engineer", "STRONG_RELATED"),
        ("Python Backend Developer using FastAPI", "Java Architect", "MISMATCH"),
    ],
)
def test_generic_role_aliases(resume, job_title, expected):
    job = SimpleNamespace(title=job_title, description=job_title, skills_required="", location="", work_mode="", job_type="")
    assert calculate_weighted_ats_score(resume, job)["role_alignment"] == expected


def test_canonical_role_ids_and_confidence():
    devops = classify_role("DevOps Engineer using Docker Kubernetes Terraform")
    assert devops["family_id"] == "DEVOPS"
    assert devops["confidence"] >= 0.6


def test_role_aware_gate_handles_partial_unknown_and_mismatch():
    assert role_aware_optimization_eligibility(69, "PARTIAL_TRANSFERABLE")["optimization_allowed"] is False
    partial = role_aware_optimization_eligibility(70, "PARTIAL_TRANSFERABLE")
    assert partial["optimization_allowed"] is True
    assert partial["eligibility_warning"]
    assert role_aware_optimization_eligibility(65, "UNKNOWN")["optimization_allowed"] is True
    assert role_aware_optimization_eligibility(90, "MISMATCH")["optimization_allowed"] is False


def test_devops_step_two_and_score_share_canonical_requirements_without_stale_cap():
    technology_list = ["AWS", "Linux", "Git", "Jenkins", "GitHub Actions", "Docker", "Kubernetes", "Helm", "Terraform", "Ansible", "Argo CD", "Python", "Bash", "Maven", "Nexus", "Splunk", "Prometheus", "Grafana", "Jira"]
    technologies = ", ".join(technology_list)
    resume = f"""Senior DevOps Engineer
Professional Experience
Acme | DevOps Engineer | 2020-Present
- Implemented and managed {technologies} for production delivery and monitoring.
"""
    job = SimpleNamespace(
        id=25, title="Senior DevOps Engineer / AWS", description=f"Requires professional use of {technologies}.",
        skills_required=technologies, location="", work_mode="", job_type="",
    )
    score = calculate_weighted_ats_score(resume, job)
    grouped = group_canonical_score_evidence(score, job.id)
    score_requirements = {item["normalized_skill"] for item in score["requirement_evidence"]}
    step_two_requirements = {item["normalized_skill"] for item in grouped["skills"]}
    assert score["role_alignment"] == "EXACT"
    assert score_requirements == step_two_requirements
    assert grouped["summary"]["total_requirements"] == sum(grouped["summary"][key] for key in ("verified", "weak_or_indirect", "missing_mandatory", "missing_preferred"))
    assert not score["missing_mandatory_skills"]
    assert score["score_cap_applied"] != 35


def test_exact_role_all_weak_skills_is_transparent_and_not_missing_skill_capped():
    resume = "DevOps Engineer\nTechnical Skills\nCI/CD, Docker, Kubernetes, Terraform, Jenkins"
    job = SimpleNamespace(
        id=26, title="DevOps Engineer", description="CI/CD Docker Kubernetes Terraform Jenkins",
        skills_required="CI/CD, Docker, Kubernetes, Terraform, Jenkins", location="", work_mode="", job_type="",
    )
    score = calculate_weighted_ats_score(resume, job)
    grouped = group_canonical_score_evidence(score, job.id)
    assert score["role_alignment"] == "EXACT"
    assert not score["missing_mandatory_skills"]
    assert grouped["summary"]["weak_or_indirect"] > 0
    assert score["score_cap_applied"] != 35
    assert score["component_details"]["mandatory_skills"]["earned"] > 0


def test_senior_devops_professional_evidence_years_or_groups_and_neutral_education():
    resume = """Jane Candidate
Senior DevOps Engineer
Technical Skills
AWS, Kubernetes, Docker, Terraform, Jenkins, GitHub Actions, GitLab CI, Linux, Python, Bash, CI/CD, Helm, Prometheus, Grafana, Ansible, Argo CD
Professional Experience
Acme | Senior DevOps Engineer | 2017-Present
- Managed Kubernetes clusters and Docker workloads across production AWS environments.
- Provisioned AWS infrastructure using Terraform and Linux automation.
- Implemented CI/CD pipelines using Jenkins, GitHub Actions, GitLab CI, Python and Bash.
- Deployed Helm and Argo CD and monitored services with Prometheus and Grafana.
"""
    job = SimpleNamespace(
        id=27, title="Senior DevOps Engineer",
        description="""Required Skills & Experience
6+ years DevOps Platform Cloud experience.
AWS, Kubernetes, Terraform, Docker, Linux, CI/CD, Python or Bash, production cloud support.
Desired Skills & Experience
Helm, ArgoCD, Datadog, Prometheus, Grafana, AWS certification, mentoring.""",
        skills_required="", location="Remote", work_mode="Remote", job_type="Full-time",
    )
    score = calculate_weighted_ats_score(resume, job)
    assert score["role_alignment"] == "EXACT"
    assert score["missing_mandatory_skills"] == []
    assert score["component_details"]["mandatory_skills"]["earned"] >= 28
    assert score["component_details"]["experience_evidence"]["earned"] > 0
    assert score["relevant_experience"]["years"] >= 6
    assert score["education_category_applicable"] is False
    assert score["component_details"]["education_certifications"]["earned"] == 5
    assert score["requirements"]["or_groups"] == [{"type": "OR_GROUP", "skills": ["Bash", "Python"]}]
    assert score["final_score"] > 46
    assert all(
        item["evidence_strength"] == "strong"
        for item in score["requirement_evidence"]
        if item["skill"] in {"AWS", "Kubernetes", "Terraform", "Docker", "Linux", "CI/CD", "Python", "Bash"}
    )


def test_date_parser_returns_full_years_not_century_prefixes():
    from app.services.resume_structure_parser import parse_resume_structure
    parsed = parse_resume_structure("Jane\nProfessional Experience\nAcme | DevOps Engineer | 2018-2024\n- Managed AWS")
    assert parsed["employments"][0]["start_year"] == 2018
    assert parsed["employments"][0]["end_year"] == 2024


def test_docx_renderer_preserves_template_and_never_overwrites_original(tmp_path):
    source = tmp_path / "original.docx"
    output = tmp_path / "optimized.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.42)
    section.left_margin = Inches(0.61)
    header = section.header.paragraphs[0]
    header.text = "Candidate Header"
    heading = document.add_paragraph()
    run = heading.add_run("PROFESSIONAL SUMMARY")
    run.bold = True
    run.font.name = "Aptos"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(20, 40, 80)
    bullet = document.add_paragraph("Original bullet", style="List Bullet")
    bullet.paragraph_format.left_indent = Inches(0.25)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Contact"
    table.cell(0, 1).text = "candidate@example.test"
    document.save(source)
    original_bytes = source.read_bytes()
    original_rendered = Document(source)

    result = render_optimized_document(str(source), "docx", "PROFESSIONAL SUMMARY\nImproved bullet\nContact\ncandidate@example.test", output_path=str(output))
    rendered = Document(output)
    assert result["template_preserved"] is True
    assert source.read_bytes() == original_bytes
    assert output != source
    assert rendered.sections[0].top_margin == original_rendered.sections[0].top_margin
    assert rendered.sections[0].left_margin == original_rendered.sections[0].left_margin
    assert rendered.sections[0].header.paragraphs[0].text == "Candidate Header"
    assert rendered.paragraphs[0].runs[0].bold is True
    assert rendered.paragraphs[0].runs[0].font.name == "Aptos"
    assert rendered.paragraphs[0].runs[0].font.size == Pt(14)
    assert rendered.paragraphs[1].style.name == "List Bullet"
    assert rendered.paragraphs[1].paragraph_format.left_indent == Inches(0.25)
    assert len(rendered.tables) == 1


def test_pdf_and_txt_renderer_are_explicitly_approximate(tmp_path):
    for file_type in ("pdf", "txt"):
        result = render_optimized_document(str(tmp_path / f"source.{file_type}"), file_type, "SUMMARY\n- Safe bullet")
        assert result["template_preserved"] is False
        assert Path(result["file_path"]).is_file()
        assert "DOCX" in result["template_note"] or "styling" in result["template_note"]


def test_edit_endpoint_saves_separate_text_and_recalculates(client, db_session, test_user, test_resume, java_job, tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph(test_resume.parsed_text)
    doc.save(source)
    test_resume.stored_file_path = str(source)
    test_resume.file_type = "docx"
    optimized = OptimizedResume(
        user_id=test_user.id, resume_id=test_resume.id, job_id=java_job.id,
        optimized_text=test_resume.parsed_text, generated_resume_text=test_resume.parsed_text,
        final_resume_text=test_resume.parsed_text, original_score=70, generated_score=72,
        template_preserved=True, source_file_type="docx", validation_status="passed",
    )
    db_session.add(optimized)
    db_session.commit()
    db_session.refresh(optimized)

    edited = test_resume.parsed_text.replace("Java", "Java and Spring Boot", 1)
    response = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": edited},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["user_modified"] is True
    assert isinstance(payload["edited_score"], int)
    db_session.refresh(optimized)
    assert optimized.generated_resume_text == test_resume.parsed_text
    assert optimized.edited_resume_text == edited
    assert optimized.final_file_path != str(source)


def test_edit_endpoint_accepts_document_and_score_only(client, db_session, test_user, test_resume, java_job, tmp_path):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph(test_resume.parsed_text)
    doc.save(source)
    test_resume.stored_file_path = str(source)
    test_resume.file_type = "docx"
    optimized = OptimizedResume(
        user_id=test_user.id, resume_id=test_resume.id, job_id=java_job.id,
        optimized_text=test_resume.parsed_text, generated_resume_text=test_resume.parsed_text,
        final_resume_text=test_resume.parsed_text, original_score=70, generated_score=72,
        template_preserved=True, source_file_type="docx", validation_status="passed",
    )
    db_session.add(optimized)
    db_session.commit()
    db_session.refresh(optimized)

    get_response = client.get(f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}")
    assert get_response.status_code == 200
    document = get_response.json()["document"]
    assert document["header"]
    assert document["sections"]

    # Score-only must not persist edits.
    document["header"]["title"] = "Temporary Title For Score"
    score_response = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "document": document, "score_only": True, "persist": False},
    )
    assert score_response.status_code == 200
    assert score_response.json()["persisted"] is False
    db_session.refresh(optimized)
    assert optimized.edited_resume_text in (None, test_resume.parsed_text)

    persist_response = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "document": document},
    )
    assert persist_response.status_code == 200
    assert persist_response.json()["document"]["header"]["title"] == "Temporary Title For Score"
    assert "score" in persist_response.json()

    db_session.refresh(optimized)
    optimized.validation_status = "passed"
    db_session.add(optimized)
    db_session.commit()

    pdf = client.get(f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith("application/pdf")


def test_edit_endpoint_rejects_cross_user_access(client, db_session, test_user, second_user, test_resume, java_job):
    optimized = OptimizedResume(user_id=test_user.id, resume_id=test_resume.id, job_id=java_job.id, optimized_text="safe")
    db_session.add(optimized)
    db_session.commit()
    response = client.patch(
        f"/api/optimizer/resume/{second_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": "changed"},
    )
    assert response.status_code == 404
