"""Regression tests for evidence-gated scoring and export integrity."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document

from app.models.optimized_resume_model import OptimizedResume
from app.models.skill_confirmation_model import UserConfirmedSkill
from app.services.ats_scoring import calculate_weighted_ats_score
from app.services.resume_generator import (
    _employment_bullet,
    _is_evidence_noise,
    _resume_ready_evidence_sentence,
)
from app.services.template_preservation_service import render_optimized_document


JAVA_JOB = SimpleNamespace(
    title="Java SDET",
    company="Acme",
    description=(
        "Looking for an SDET with Java, Selenium, JMeter, Appium, Cucumber, "
        "and TestNG experience for API and UI automation."
    ),
    skills_required="Java, Selenium, JMeter, Appium, Cucumber, TestNG",
)


BASE_RESUME = """
Jane Developer
jane@example.com | 555-111-2222

Professional Summary
Java SDET with Selenium and TestNG automation experience.

Technical Skills
Languages: Java
Testing: Selenium, TestNG
Tools: Jenkins, Git

Professional Experience
Acme Corp | SDET | 2020 - Present
- Built Java Selenium suites for regression testing.
- Automated API checks with TestNG.

Education
B.S. Computer Science
""".strip()


def test_rejected_appium_typed_gets_zero_credit():
    with_appium = BASE_RESUME + "\nTechnical Skills\nJava, Selenium, Appium, Appium, Appium"
    denied = ["Appium"]
    baseline = calculate_weighted_ats_score(
        BASE_RESUME, JAVA_JOB,
        original_resume_text=BASE_RESUME,
        denied_skills=denied,
        evidence_gated=True,
    )
    boosted_attempt = calculate_weighted_ats_score(
        with_appium, JAVA_JOB,
        original_resume_text=BASE_RESUME,
        denied_skills=denied,
        evidence_gated=True,
        confirmed_skills={},
    )
    assert boosted_attempt["match_percentage"] <= baseline["match_percentage"]
    assert "Appium" in (boosted_attempt.get("rejected_present_requirements") or [])
    assert boosted_attempt["match_percentage"] == baseline["match_percentage"]


def test_rejected_cucumber_typed_gets_zero_credit():
    with_cuke = BASE_RESUME.replace("TestNG", "TestNG, Cucumber")
    denied = ["Cucumber"]
    baseline = calculate_weighted_ats_score(
        BASE_RESUME, JAVA_JOB,
        original_resume_text=BASE_RESUME,
        denied_skills=denied,
        evidence_gated=True,
    )
    after = calculate_weighted_ats_score(
        with_cuke, JAVA_JOB,
        original_resume_text=BASE_RESUME,
        denied_skills=denied,
        evidence_gated=True,
    )
    assert after["match_percentage"] == baseline["match_percentage"]
    assert "Cucumber" in (after.get("rejected_present_requirements") or [])


def test_confirmed_jmeter_present_then_removed():
    confirmed = {
        "JMeter": {
            "credit": 0.85,
            "evidence_status": "USER_CONFIRMED_PROFESSIONAL",
            "evidence_type": "Current employment",
            "usage_description": "Used JMeter for API load testing",
        }
    }
    with_jmeter = BASE_RESUME + "\n- Performed load testing with JMeter for concurrent API traffic."
    present = calculate_weighted_ats_score(
        with_jmeter, JAVA_JOB,
        confirmed_skills=confirmed,
        original_resume_text=BASE_RESUME,
        denied_skills=["Appium", "Cucumber"],
        evidence_gated=True,
    )
    removed = calculate_weighted_ats_score(
        BASE_RESUME, JAVA_JOB,
        confirmed_skills=confirmed,
        original_resume_text=BASE_RESUME,
        denied_skills=["Appium", "Cucumber"],
        evidence_gated=True,
    )
    assert present["match_percentage"] >= removed["match_percentage"]
    assert "JMeter" in (present.get("confirmed_skills_used") or present.get("confirmed_present_requirements") or [])
    assert "JMeter" not in (removed.get("confirmed_skills_used") or [])


def test_unknown_manual_skill_gets_no_credit():
    with_playwright = BASE_RESUME + "\nSkills: Playwright"
    # Playwright is not on original and not confirmed — even if JD mentioned it wouldn't apply here;
    # force it as preferred-like by ensuring evidence_gated zeroes unverified JD skills only.
    # For JD skills not in this resume originally: Selenium stays; inventing Appium already covered.
    score = calculate_weighted_ats_score(
        with_playwright, JAVA_JOB,
        original_resume_text=BASE_RESUME,
        evidence_gated=True,
    )
    assert "Playwright" not in (score.get("confirmed_skills_used") or [])
    # Unverified list only includes JD requirements that are present but ungated.
    assert score["match_percentage"] >= 0


def test_evidence_noise_filtered_and_resume_ready():
    assert _is_evidence_noise("Client: status meetings.")
    # A textbook-definition answer ("JMeter is mainly used for...") is
    # SKILL_KNOWLEDGE, not PROFESSIONAL_USAGE_EVIDENCE — it must never be
    # salvaged into a manufactured bullet. (Was previously polished into a
    # generic "Performed performance and load testing..." bullet regardless
    # of what the skill actually was; that fallback has been removed.)
    polished = _resume_ready_evidence_sentence(
        "JMeter",
        "Apache JMeter is mainly used for performance and load testing. It helps us check APIs.",
    )
    assert polished == ""
    bullet = _employment_bullet({
        "skill": "JMeter",
        "usage_description": "Client: status meetings.",
        "evidence_status": "USER_CONFIRMED_PROFESSIONAL",
    })
    # Noise-only description produces no manufactured Professional Experience
    # bullet at all — the skill still reaches the resume via Technical
    # Skills placement (a separate mechanism), never a fabricated claim.
    assert bullet == ""


def test_docx_template_preserves_heading_and_table(tmp_path):
    source = tmp_path / "original.docx"
    doc = Document()
    title = doc.add_paragraph("JANE DEVELOPER")
    title.runs[0].bold = True
    heading = doc.add_paragraph("TECHNICAL SKILLS")
    heading.runs[0].bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Java"
    table.cell(1, 0).text = "Testing"
    table.cell(1, 1).text = "Selenium, TestNG"
    doc.add_paragraph("PROFESSIONAL EXPERIENCE")
    doc.add_paragraph("Acme Corp | SDET | 2020 - Present")
    doc.add_paragraph("Built Java Selenium suites.")
    doc.save(source)

    optimized = """
Jane Developer

Technical Skills
Languages: Java
Testing: Selenium, TestNG, JMeter

Professional Experience
Acme Corp | SDET | 2020 - Present
- Built Java Selenium suites.
- Performed load testing with JMeter.
""".strip()
    rendered = render_optimized_document(str(source), "docx", optimized, output_path=str(tmp_path / "out.docx"))
    assert rendered["template_preserved"] is True
    out = Document(rendered["file_path"])
    texts = [p.text for p in out.paragraphs]
    assert any(t == "TECHNICAL SKILLS" for t in texts)
    assert out.tables, "skills table must be preserved"
    assert "Java" in out.tables[0].cell(0, 1).text


def test_recalculate_endpoint_rejects_appium_credit(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    source = tmp_path / "source.docx"
    Document().save(source)
    test_resume.stored_file_path = str(source)
    test_resume.file_type = "docx"
    test_resume.parsed_text = BASE_RESUME
    java_job.description = (
        "SDET role requiring Java, Selenium, JMeter, Appium, Cucumber, and TestNG."
    )
    java_job.skills_required = "Java, Selenium, JMeter, Appium, Cucumber, TestNG"
    db_session.add(test_resume)
    db_session.add(java_job)
    optimized = OptimizedResume(
        user_id=test_user.id, resume_id=test_resume.id, job_id=java_job.id,
        optimized_text=BASE_RESUME, generated_resume_text=BASE_RESUME,
        final_resume_text=BASE_RESUME, original_score=70, generated_score=75,
        edited_score=75, validation_status="passed",
    )
    db_session.add(optimized)
    db_session.add(UserConfirmedSkill(
        user_id=test_user.id, job_id=java_job.id, resume_id=test_resume.id,
        skill_name="Appium", normalized_skill_name="appium",
        evidence_status="REJECTED", evidence_type=None, evidence_source="user_input",
        is_active=True, confirmed=False,
    ))
    db_session.commit()
    db_session.refresh(optimized)

    baseline = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={
            "optimized_resume_id": optimized.id,
            "edited_resume_text": BASE_RESUME,
            "action": "recalculate",
            "persist": True,
        },
    ).json()["edited_score"]

    after = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={
            "optimized_resume_id": optimized.id,
            "edited_resume_text": BASE_RESUME + "\nAppium Appium Appium",
            "action": "recalculate",
            "persist": True,
        },
    ).json()
    assert after["edited_score"] <= baseline
    warnings = after.get("score", {}).get("skill_credit_warnings") or []
    assert any("Appium" in w and "No ATS credit" in w for w in warnings)


def test_pdf_export_uses_docx_conversion(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Jane Developer")
    doc.add_paragraph("Professional Summary")
    doc.add_paragraph("Java SDET")
    doc.save(source)
    test_resume.stored_file_path = str(source)
    test_resume.file_type = "docx"
    test_resume.parsed_text = "Jane Developer\nProfessional Summary\nJava SDET"
    db_session.add(test_resume)
    optimized = OptimizedResume(
        user_id=test_user.id, resume_id=test_resume.id, job_id=java_job.id,
        optimized_text=test_resume.parsed_text, generated_resume_text=test_resume.parsed_text,
        final_resume_text=test_resume.parsed_text, validation_status="passed",
        original_score=70, generated_score=72,
    )
    db_session.add(optimized)
    db_session.commit()
    db_session.refresh(optimized)

    fake_pdf = tmp_path / "converted.pdf"
    fake_pdf.write_bytes(b"%PDF-1.4\n" + b"0" * 600)

    def _fake_convert(docx_path, destination=None):
        target = Path(destination) if destination else fake_pdf
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(fake_pdf.read_bytes())
        return target

    with patch("app.routes.optimizer_edit_routes.convert_docx_to_pdf", side_effect=_fake_convert) as convert:
        with patch("app.routes.optimizer_edit_routes.validate_pdf_export", return_value={"valid": True, "page_count": 1}):
            response = client.get(
                f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=pdf"
            )
    assert response.status_code == 200
    assert convert.called
    assert response.headers["content-type"].startswith("application/pdf")
