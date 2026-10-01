"""Editor/export document-version consistency guard.

Verifies the minimal, no-new-DB-column guarantee: the resume content shown
in the Step 5 editor at download time is the exact content source used for
export, and a background/retry regeneration can never silently overwrite a
user's already-saved edits. No OpenAI/ATS-scoring/evidence-rule change is
involved — `document_version` is a pure content hash of the same
`final_resume_text` field that already backs every read/write path, and the
actual protection is the `user_modified` gate in
`ats_optimization_workflow_service.execute_workflow`.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document

from app.models.optimized_resume_model import OptimizedResume
from app.services.ats_optimization_workflow_service import execute_workflow, start_or_get_workflow


def _seed_optimized(db_session, test_user, test_resume, java_job, text=None, validation_status="passed"):
    body = text or test_resume.parsed_text
    optimized = OptimizedResume(
        user_id=test_user.id,
        resume_id=test_resume.id,
        job_id=java_job.id,
        optimized_text=body,
        generated_resume_text=body,
        final_resume_text=body,
        original_score=70,
        generated_score=80,
        edited_score=80,
        template_preserved=True,
        source_file_type="docx",
        validation_status=validation_status,
    )
    db_session.add(optimized)
    db_session.commit()
    db_session.refresh(optimized)
    return optimized


def _with_docx_source(db_session, test_resume, tmp_path):
    source = tmp_path / "source.docx"
    Document().save(source)
    test_resume.stored_file_path = str(source)
    test_resume.file_type = "docx"
    db_session.add(test_resume)
    db_session.commit()


def test_save_returns_a_document_version_that_tracks_content(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(db_session, test_user, test_resume, java_job)

    text_a = test_resume.parsed_text + "\nAdditional marker A."
    resp_a = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": text_a, "action": "save", "persist": True},
    )
    assert resp_a.status_code == 200
    version_a = resp_a.json()["document_version"]
    assert version_a

    text_b = test_resume.parsed_text + "\nAdditional marker B."
    resp_b = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": text_b, "action": "save", "persist": True},
    )
    version_b = resp_b.json()["document_version"]
    assert version_b != version_a, "different saved content must produce a different document_version"

    resp_a_again = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": text_a, "action": "save", "persist": True},
    )
    assert resp_a_again.json()["document_version"] == version_a, (
        "saving the same content again must reproduce the same document_version"
    )


def test_recalculate_document_version_matches_a_subsequent_get(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(db_session, test_user, test_resume, java_job)

    marker = test_resume.parsed_text + "\nRecalculate marker."
    save_resp = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={
            "optimized_resume_id": optimized.id,
            "edited_resume_text": marker,
            "action": "recalculate",
            "persist": True,
        },
    )
    assert save_resp.status_code == 200
    recalculate_version = save_resp.json()["document_version"]

    get_resp = client.get(f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}")
    assert get_resp.status_code == 200
    assert get_resp.json()["document_version"] == recalculate_version


def test_download_carries_x_document_version_header_matching_the_save(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(db_session, test_user, test_resume, java_job)

    # Edit the (mutable) summary — appending after Certifications instead
    # would trip the unrelated protected-section truthfulness check and
    # fail validation, which isn't what this test is about.
    marker = test_resume.parsed_text.replace(
        "Java Full Stack Developer with experience building enterprise applications.",
        "Download version marker. Java Full Stack Developer with experience building enterprise applications.",
    )
    save_resp = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": marker, "action": "save", "persist": True},
    )
    saved_version = save_resp.json()["document_version"]

    docx_resp = client.get(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=docx"
    )
    assert docx_resp.status_code == 200
    assert docx_resp.headers.get("x-document-version") == saved_version


def test_pdf_download_x_document_version_matches_docx_for_the_same_content(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    """The PDF is generated from the exact same rendered DOCX/version as the
    DOCX download for identical content — proving the PDF isn't served from
    some other, potentially stale, generation snapshot."""
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(
        db_session, test_user, test_resume, java_job, validation_status="passed_with_warnings",
    )

    docx_resp = client.get(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=docx"
    )
    docx_version = docx_resp.headers.get("x-document-version")
    assert docx_version

    fake_pdf = tmp_path / "converted.pdf"
    fake_pdf.write_bytes(b"%PDF-1.4\n" + b"0" * 600)

    def _fake_convert(docx_path, destination=None):
        target = Path(destination) if destination else fake_pdf
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(fake_pdf.read_bytes())
        return target

    with patch("app.routes.optimizer_edit_routes.convert_docx_to_pdf", side_effect=_fake_convert):
        with patch("app.routes.optimizer_edit_routes.validate_pdf_export", return_value={"valid": True, "page_count": 1}):
            pdf_resp = client.get(
                f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=pdf"
            )
    assert pdf_resp.status_code == 200
    assert pdf_resp.headers.get("x-document-version") == docx_version


def test_editing_summary_then_saving_then_downloading_shows_the_new_summary(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(db_session, test_user, test_resume, java_job)

    new_summary = "UNIQUE_SUMMARY_MARKER Backend engineer specializing in distributed systems."
    edited = test_resume.parsed_text.replace(
        "Java Full Stack Developer with experience building enterprise applications.",
        new_summary,
    )
    assert edited != test_resume.parsed_text

    save_resp = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": edited, "action": "save", "persist": True},
    )
    assert save_resp.status_code == 200

    download_resp = client.get(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=docx"
    )
    assert download_resp.status_code == 200
    loaded = Document(BytesIO(download_resp.content))
    text = "\n".join(p.text for p in loaded.paragraphs)
    assert "UNIQUE_SUMMARY_MARKER" in text


def test_editing_a_bullet_then_saving_then_downloading_shows_the_new_bullet(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(db_session, test_user, test_resume, java_job)

    new_bullet = "Led a UNIQUE_BULLET_MARKER migration to a service-mesh architecture."
    edited = test_resume.parsed_text.replace(
        "- Built React components for customer portals.",
        f"- {new_bullet}",
    )
    assert edited != test_resume.parsed_text

    save_resp = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": edited, "action": "save", "persist": True},
    )
    assert save_resp.status_code == 200

    download_resp = client.get(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}/{optimized.id}/download?format=docx"
    )
    assert download_resp.status_code == 200
    loaded = Document(BytesIO(download_resp.content))
    text = "\n".join(p.text for p in loaded.paragraphs)
    assert "UNIQUE_BULLET_MARKER" in text


def test_workflow_regeneration_cannot_overwrite_a_saved_editor_edit(
    db_session, test_user, test_resume, java_job, tmp_path,
):
    """The exact scenario the version guard exists to prevent: a background/
    retry regeneration (execute_workflow) must never clobber final_resume_text
    once the user has saved real edits to it — only the informational
    'generated' snapshot may keep refreshing."""
    _with_docx_source(db_session, test_resume, tmp_path)

    workflow = start_or_get_workflow(db_session, user_id=test_user.id, job_id=java_job.id)
    execute_workflow(db_session, workflow.id)

    optimized = (
        db_session.query(OptimizedResume)
        .filter(
            OptimizedResume.user_id == test_user.id,
            OptimizedResume.resume_id == test_resume.id,
            OptimizedResume.job_id == java_job.id,
        )
        .order_by(OptimizedResume.id.desc())
        .first()
    )
    assert optimized is not None
    first_generated_score = optimized.generated_score

    user_saved_text = "UNIQUE_USER_SAVED_TEXT_THAT_MUST_SURVIVE_A_REGENERATION\n" + (test_resume.parsed_text or "")
    optimized.final_resume_text = user_saved_text
    optimized.edited_resume_text = user_saved_text
    optimized.user_modified = True
    db_session.add(optimized)
    db_session.commit()

    forced = start_or_get_workflow(db_session, user_id=test_user.id, job_id=java_job.id, force=True)
    assert forced.id != workflow.id
    execute_workflow(db_session, forced.id)

    db_session.refresh(optimized)
    assert optimized.final_resume_text == user_saved_text, (
        "a regeneration overwrote the user's saved edit — the exact regression this guard prevents"
    )
    assert optimized.user_modified is True
    # The informational "latest generation" snapshot still refreshes.
    assert optimized.generated_resume_text is not None
    assert optimized.generated_score is not None


def test_protected_employer_metadata_unchanged_after_a_summary_only_save(
    client, db_session, test_user, test_resume, java_job, tmp_path,
):
    _with_docx_source(db_session, test_resume, tmp_path)
    optimized = _seed_optimized(db_session, test_user, test_resume, java_job)

    edited = test_resume.parsed_text.replace(
        "Java Full Stack Developer with experience building enterprise applications.",
        "UNIQUE_SUMMARY_ONLY_EDIT Backend engineer.",
    )
    save_resp = client.patch(
        f"/api/optimizer/resume/{test_user.id}/{java_job.id}",
        json={"optimized_resume_id": optimized.id, "edited_resume_text": edited, "action": "save", "persist": True},
    )
    assert save_resp.status_code == 200
    saved_text = save_resp.json()["final_text"]

    assert "Acme Corp | Java Full Stack Developer | Jan 2018 - Present" in saved_text
    assert "Developed REST APIs using Java and Spring Boot for account workflows." in saved_text
    assert "Built React components for customer portals." in saved_text
    assert "Deployed services with Docker and Jenkins on AWS." in saved_text
